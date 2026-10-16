"""Génère les modèles Word téléchargeables de la page Ressources (CV étudiant, lettre de motivation).

Usage : python3 web/scripts/modeles_docx.py  →  web/public/ressources/modeles/*.docx
"""
import os
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "public", "ressources", "modeles")
BLEU = RGBColor(0x1A, 0x47, 0xF5)
ENCRE = RGBColor(0x0B, 0x15, 0x33)
GRIS = RGBColor(0x6B, 0x74, 0x93)


def base(titre):
    d = Document()
    d.core_properties.title = titre
    d.core_properties.author = "Navigoal"
    for s in d.sections:
        s.top_margin = s.bottom_margin = Cm(1.6)
        s.left_margin = s.right_margin = Cm(1.8)
    st = d.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)
    st.font.color.rgb = ENCRE
    st.paragraph_format.space_after = Pt(3)
    return d


def run(p, texte, gras=False, couleur=None, taille=None, italique=False):
    r = p.add_run(texte)
    r.bold = gras
    r.italic = italique
    if couleur:
        r.font.color.rgb = couleur
    if taille:
        r.font.size = Pt(taille)
    return r


def filet(p):
    pPr = p._p.get_or_add_pPr()
    b = OxmlElement("w:pBdr")
    bot = OxmlElement("w:bottom")
    for k, v in {"w:val": "single", "w:sz": "8", "w:space": "1", "w:color": "1A47F5"}.items():
        bot.set(qn(k), v)
    b.append(bot)
    pPr.append(b)


def rubrique(d, titre):
    p = d.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    run(p, titre.upper(), gras=True, couleur=BLEU, taille=11)
    filet(p)


def ligne(d, gauche, droite="", detail=None):
    p = d.add_paragraph()
    p.paragraph_format.tab_stops.add_tab_stop(Cm(17.4), alignment=2)
    run(p, gauche, gras=True)
    if droite:
        run(p, "\t" + droite, couleur=GRIS)
    for t in detail or []:
        q = d.add_paragraph(style="List Bullet")
        run(q, t)


def cv():
    d = base("Modèle de CV étudiant — Navigoal")
    p = d.add_paragraph()
    run(p, "Prénom NOM", gras=True, taille=22)
    p = d.add_paragraph()
    run(p, "Candidat·e en [formation visée] · [Ville], Maroc", couleur=BLEU, taille=12)
    p = d.add_paragraph()
    run(p, "Téléphone (avec indicatif) · adresse e-mail professionnelle · ville et pays · nationalité · lien LinkedIn (facultatif)", couleur=GRIS, taille=9.5)
    rubrique(d, "Profil")
    d.add_paragraph("En 2 ou 3 lignes : qui tu es, ce que tu vises et ce que tu apportes. Exemple : « Bachelier série scientifique, passionné par les données, je souhaite intégrer un cycle d'ingénieur en informatique pour devenir data analyst. »")
    rubrique(d, "Formation")
    ligne(d, "Baccalauréat [série] — mention [mention]", "Année", ["Établissement, ville, pays", "Moyenne générale : xx/20 (ajoute les matières fortes)"])
    ligne(d, "Classe de première / seconde", "Années", ["Établissement, ville, pays"])
    rubrique(d, "Expériences et projets")
    ligne(d, "Stage, job d'été, bénévolat ou projet scolaire", "Mois année", ["Ce que tu as fait (verbe d'action + résultat)", "Ce que tu as appris"])
    ligne(d, "Projet personnel (site web, club, association…)", "Année", ["Ton rôle et le résultat obtenu"])
    rubrique(d, "Compétences")
    d.add_paragraph("Informatique : bureautique, outils numériques, langages éventuels", style="List Bullet")
    d.add_paragraph("Qualités : rigueur, travail en équipe, autonomie (illustre-les dans tes expériences)", style="List Bullet")
    rubrique(d, "Langues")
    d.add_paragraph("Français : courant · Anglais : niveau (A2, B1…) · Arabe ou langue nationale : niveau", style="List Bullet")
    rubrique(d, "Centres d'intérêt")
    d.add_paragraph("2 ou 3 activités précises (sport en club, lecture, musique, engagement associatif).")
    p = d.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    run(p, "Conseils : une seule page, pas de photo obligatoire, police lisible, enregistre en PDF sous « NOM_Prenom_CV.pdf » avant d'envoyer. Modèle Navigoal — navigoal.com/ressources", italique=True, couleur=GRIS, taille=8.5)
    d.save(os.path.join(OUT, "modele-cv-etudiant-navigoal.docx"))


def lettre():
    d = base("Modèle de lettre de motivation — Navigoal")
    for t in ["Prénom NOM", "Adresse complète, ville, pays", "Téléphone (avec indicatif) · e-mail"]:
        p = d.add_paragraph()
        run(p, t, gras=t == "Prénom NOM")
    p = d.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run(p, "Nom de l'établissement\nService des admissions\nAdresse, ville, Maroc")
    p = d.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run(p, "[Ville], le [date]")
    p = d.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    run(p, "Objet : candidature en [intitulé exact de la formation] — rentrée [année]", gras=True)
    d.add_paragraph("Madame, Monsieur,")
    blocs = [
        ("Vous", "Explique pourquoi cet établissement : son label (reconnu par l'État, filière accréditée), son programme, ses stages, son campus de [ville], ses partenariats. Montre que tu connais l'école."),
        ("Moi", "Présente ton parcours et tes atouts en lien avec la formation : série du bac, matières fortes, projets, expériences, qualités illustrées par un exemple concret."),
        ("Nous", "Décris ton projet professionnel (métier visé) et ce que tu apporteras à la promotion. Mentionne ta motivation pour étudier au Maroc et ta préparation (visa, logement, budget)."),
    ]
    for titre, texte in blocs:
        p = d.add_paragraph()
        run(p, f"[{titre}] ", gras=True, couleur=BLEU)
        run(p, texte)
    d.add_paragraph("Je me tiens à votre disposition pour un entretien, en présentiel ou à distance, et vous prie d'agréer, Madame, Monsieur, l'expression de mes salutations distinguées.")
    p = d.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run(p, "Prénom NOM\n(signature)")
    p = d.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    run(p, "Conseils : une page maximum, une lettre personnalisée par établissement, pas de copier-coller, relis l'orthographe, enregistre en PDF. Modèle Navigoal — navigoal.com/ressources", italique=True, couleur=GRIS, taille=8.5)
    d.save(os.path.join(OUT, "modele-lettre-motivation-navigoal.docx"))


if __name__ == "__main__":
    os.makedirs(OUT, exist_ok=True)
    cv()
    lettre()
    print("Modèles générés dans", os.path.normpath(OUT))
